# modules/loaders.py
import os
import io
import pdfplumber
import pandas as pd
from docx import Document
from PIL import Image

def load_document(file_path: str) -> dict:
    """
    Ingests multi-format logistics documents and standardizes content output.
    
    Returns:
        dict: {
            "file_path": str,
            "file_name": str,
            "file_type": str,
            "text": str,
            "images": list[Image.Image]  # List of PIL Image objects
        }
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Document not found at: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    file_name = os.path.basename(file_path)

    result = {
        "file_path": file_path,
        "file_name": file_name,
        "file_type": ext,
        "text": "",
        "images": []
    }

    # 1. PDF Files
    if ext == ".pdf":
        extracted_text = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                extracted_text.append(txt)
                img = page.to_image(resolution=150).original
                result["images"].append(img)
        result["text"] = "\n".join(extracted_text)

    # 2. Word Documents (.docx) - Includes Embedded Image Extraction
    elif ext == ".docx":
        doc = Document(file_path)
        full_text = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                full_text.append(" | ".join([cell.text.strip() for cell in row.cells]))
        result["text"] = "\n".join(full_text)

        # Extract embedded images for OCR fallback when text is empty/scanned
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    img_bytes = rel.target_part.blob
                    img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
                    result["images"].append(img)
                except Exception:
                    pass

    # 3. Excel Spreadsheets (.xlsx)
    elif ext == ".xlsx":
        df_dict = pd.read_excel(file_path, sheet_name=None)
        sheets_text = []
        for sheet_name, df in df_dict.items():
            sheets_text.append(f"--- Sheet: {sheet_name} ---")
            sheets_text.append(df.to_string(index=False))
        result["text"] = "\n".join(sheets_text)

    # 4. Plain Text (.txt)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            result["text"] = f.read()

    # 5. Standalone Image Files (.png, .jpg, .jpeg)
    elif ext in [".png", ".jpg", ".jpeg"]:
        img = Image.open(file_path).convert("RGB")
        result["images"].append(img)

    else:
        raise ValueError(f"Unsupported file extension: {ext}")

    return result